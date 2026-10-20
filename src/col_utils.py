"""Helpers for locating template columns by their header text.

The configured column letters (po_col, cost_center_col ...) are treated as
defaults only: when the header row shows the column somewhere else, the column
found by its label wins, so templates can be re-laid-out without code changes.
"""

import re

from openpyxl.utils import get_column_letter


def norm(val) -> str:
    """Lower-case, collapse whitespace/line breaks, drop '*' for header matching."""
    if val is None:
        return ""
    return re.sub(r"\s+", " ", str(val).replace("*", " ")).strip().lower()


def find_header_row(ws, default: int, max_scan: int = 60) -> int:
    """Row holding the 'Contact for PO/ER' label; ``default`` when absent."""
    max_col = min(ws.max_column or 50, 60)
    for r in range(1, min(ws.max_row or max_scan, max_scan) + 1):
        for c in range(1, max_col + 1):
            if "contact for po" in norm(ws.cell(row=r, column=c).value):
                return r
    return default


def find_po_col(ws, header_row: int, default_letter: str) -> str:
    """Column letter of the PO / ER number header; ``default_letter`` when the
    configured column already looks right or nothing better is found."""
    max_col = min(ws.max_column or 50, 80)

    def is_po(text: str) -> bool:
        return (
            text in ("po", "po #", "po#", "po number", "po no")
            or "po number" in text
            or "po/er" in text
            or "er number" in text
        )

    from openpyxl.utils import column_index_from_string
    try:
        cfg_text = norm(ws.cell(row=header_row, column=column_index_from_string(default_letter)).value)
    except ValueError:
        cfg_text = ""
    if is_po(cfg_text):
        return default_letter
    for c in range(1, max_col + 1):
        if is_po(norm(ws.cell(row=header_row, column=c).value)):
            return get_column_letter(c)
    return default_letter


def find_label_col(ws, label: str, default_letter: str, max_row: int, max_col: int = 15) -> str:
    """Column letter of the first cell above ``max_row`` whose text contains
    ``label``; ``default_letter`` when not found."""
    label = label.lower()
    for r in range(1, max(max_row, 2)):
        for c in range(1, max_col + 1):
            if label in norm(ws.cell(row=r, column=c).value):
                return get_column_letter(c)
    return default_letter
