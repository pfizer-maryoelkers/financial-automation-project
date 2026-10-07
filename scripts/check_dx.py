import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter

wb = load_workbook("data/templates/P3 Testing.xlsx")
tab = "GSC EMEA- Dx Sub-cluster \u2013 resh"
ws = wb[tab]

print(f"Tab: {tab}")
print(f"Previous Period Invoices row:", end=" ")
for r in range(13, 45):
    if ws.cell(row=r, column=1).value == "Previous Period Invoices":
        print(r)
        break

print("\nRow 9 K9 formula (first 200 chars):")
v = ws["K9"].value
if isinstance(v, ArrayFormula):
    print(" ", v.text[:200])
else:
    print(" ", repr(v))

print("\nRow 11 BS11:", ws["BS11"].value)

# Show all non-empty data rows
print("\nData rows (cols A-N only):")
for r in range(13, 40):
    vals = []
    for col in "ABCDEFGHIJKLMN":
        v = ws[f"{col}{r}"].value
        if v is not None:
            vals.append(f"{col}={repr(v)[:20]}")
    if vals:
        print(f"  R{r}: {' '.join(vals)}")
