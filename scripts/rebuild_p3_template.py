"""rebuild_p3_template.py

Rebuilds data/templates/P3 Testing.xlsx so it contains exactly two sheets
in the correct order:
  Tab 1 — "Insert P3 HERE"   (the P3 ID config sheet)
  Tab 2 — "Project Template" (the blank template to clone per P3 ID)

All old output tabs (previous pipeline run) are removed.
The blank template is taken from the current "Sheet1".
"""

import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from copy import copy
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter

SRC = "data/templates/P3 Testing.xlsx"
DST = "data/templates/P3 Testing.xlsx"

KEEP_SHEETS = {"Sheet1", "Insert P3 HERE"}
TEMPLATE_NEW_NAME = "Project Template"

wb_src = load_workbook(SRC)

# ── Verify the sheets we need exist ──────────────────────────────────────────
missing = KEEP_SHEETS - set(wb_src.sheetnames)
if missing:
    print(f"ERROR: sheet(s) not found: {missing}")
    sys.exit(1)

# ── Build a new workbook with only the two sheets in the right order ─────────
wb_new = Workbook()
# Remove the default empty sheet openpyxl creates
wb_new.remove(wb_new.active)


def _copy_sheet(src_wb, src_name, dst_wb, dst_name):
    """Copy a worksheet from src_wb into dst_wb under dst_name."""
    src_ws = src_wb[src_name]
    dst_ws = dst_wb.create_sheet(dst_name)

    # Copy cell values, data_type, number_format, and basic font/fill/border
    for row in src_ws.iter_rows():
        for cell in row:
            dst_cell = dst_ws.cell(row=cell.row, column=cell.column)
            dst_cell.value = cell.value
            dst_cell.data_type = cell.data_type
            if cell.has_style:
                dst_cell.font      = copy(cell.font)
                dst_cell.fill      = copy(cell.fill)
                dst_cell.border    = copy(cell.border)
                dst_cell.alignment = copy(cell.alignment)
                dst_cell.number_format = cell.number_format

    # Copy column widths
    for col_letter, col_dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[col_letter].width = col_dim.width

    # Copy row heights
    for row_idx, row_dim in src_ws.row_dimensions.items():
        dst_ws.row_dimensions[row_idx].height = row_dim.height

    # Copy merged cells
    for merge in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(merge))

    # Copy freeze panes
    if src_ws.freeze_panes:
        dst_ws.freeze_panes = src_ws.freeze_panes

    print(f"  Copied '{src_name}' -> '{dst_name}'")
    return dst_ws


# Tab 1: Insert P3 HERE
_copy_sheet(wb_src, "Insert P3 HERE", wb_new, "Insert P3 HERE")

# Tab 2: blank project template (from Sheet1), renamed to "Project Template"
_copy_sheet(wb_src, "Sheet1", wb_new, TEMPLATE_NEW_NAME)

# Save
wb_new.save(DST)
print(f"\nSaved clean template to: {DST}")
print(f"Sheets: {wb_new.sheetnames}")
