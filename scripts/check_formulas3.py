import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter

wb = load_workbook("data/templates/P3 Testing.xlsx")
print("All tabs:", wb.sheetnames)

# Find the GSC EMEA Dx tab
tab = next(s for s in wb.sheetnames if "Dx Sub" in s or "EMEA" in s and "Digit" not in s)
print(f"\nChecking tab: {tab}")
ws = wb[tab]

print("\nRow 9 formulas (P3 Actuals summary):")
for col in range(1, 75):
    v = ws.cell(row=9, column=col).value
    if v is not None:
        col_letter = get_column_letter(col)
        if isinstance(v, ArrayFormula):
            print(f"  {col_letter}9 ARRAY: {v.text[:120]}")
        else:
            print(f"  {col_letter}9: {repr(v)[:80]}")

print("\nRow 11 BS:", ws["BS11"].value)
print("Previous Period Invoices row:")
for r in range(13, 45):
    if ws.cell(row=r, column=1).value == "Previous Period Invoices":
        print(f"  Found at row {r}")
        break
