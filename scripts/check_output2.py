import sys, io, os
sys.path.insert(0, os.getcwd())
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

wb = load_workbook("data/templates/P3 Testing.xlsx", data_only=True)

for tab in ["GL&NS- L-Seagen Products to be", "Israel - New LSP set up - SLE", "2026 GSC GLNS Data Quality Auto"]:
    if tab not in wb.sheetnames:
        print(f"Tab not found: {tab}")
        continue
    ws = wb[tab]
    print(f"\n=== {tab} ===")
    print("All non-empty cells rows 12-20:")
    for r in range(12, 21):
        for col in range(1, 75):
            v = ws.cell(row=r, column=col).value
            if v is not None:
                print(f"  R{r} {get_column_letter(col)}: {repr(v)}")
