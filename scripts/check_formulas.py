import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

wb = load_workbook("data/templates/P3 Testing.xlsx")
# use the first P3 tab (tab index 2 = index 2 in sheetnames)
tab = wb.sheetnames[2]
print(f"Checking tab: {tab}")
ws = wb[tab]

for cell_ref in ["K9", "P9", "U9"]:
    v = ws[cell_ref].value
    if isinstance(v, ArrayFormula):
        print(f"  {cell_ref} ArrayFormula: {v.text[:120]}")
    else:
        print(f"  {cell_ref}: {repr(v)}")

print()
print(f"  BS11 (Total SUM): {ws['BS11'].value}")
print(f"  K11 (Jan date):   {ws['K11'].value}")
print(f"  BS12 header:      {ws['BS12'].value}")
