import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

wb = load_workbook("data/templates/P3 Testing.xlsx", data_only=True)
ws = wb["GSC EMEA- Dx Sub-cluster \u2013 resh"]

print("All written values (rows 13-36):")
for r in range(13, 37):
    row_data = []
    for col in range(1, 75):
        v = ws.cell(row=r, column=col).value
        if v is not None:
            row_data.append(f"{get_column_letter(col)}={v}")
    if row_data:
        print(f"  R{r}: " + "  ".join(row_data))

print()
print("Summary R9 (P3 Actuals) cached values:")
for col_letter in ["K","P","U","Z","AE","AJ","AO","AT","AY","BD","BI","BN"]:
    v = ws[f"{col_letter}9"].value
    print(f"  {col_letter}9: {v}")

print()
print("Summary R10 (Forecast Actuals) cached values:")
for col_letter in ["K","P","U","Z","AE","AJ","AO","AT","AY","BD","BI","BN"]:
    v = ws[f"{col_letter}10"].value
    print(f"  {col_letter}10: {v}")
