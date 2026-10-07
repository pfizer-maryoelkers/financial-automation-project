import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

wb = load_workbook("data/templates/P3 Testing.xlsx")
tab = "GSC EMEA- Dx Sub-cluster \u2013 resh"
ws = wb[tab]

print(f"Tab: {tab}")
print("\nSummary row 9 (P3 Actuals) - all non-empty:")
for col in range(1, 75):
    v = ws.cell(row=9, column=col).value
    if v is not None:
        print(f"  {get_column_letter(col)}9 = {v}")

print("\nSummary row 10 (Forecast Actuals) - all non-empty:")
for col in range(1, 75):
    v = ws.cell(row=10, column=col).value
    if v is not None:
        print(f"  {get_column_letter(col)}10 = {v}")

print("\nData rows monthly values (cols K-AO):")
for r in range(13, 36):
    vals = []
    for col in range(11, 42):  # K to AO
        v = ws.cell(row=r, column=col).value
        if v is not None and v != 0:
            vals.append(f"{get_column_letter(col)}={v:.0f}")
    if vals:
        print(f"  R{r}: {' '.join(vals)}")
